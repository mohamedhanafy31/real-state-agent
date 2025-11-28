#!/usr/bin/env python3
"""
Script to verify company name replacement in Word documents.
"""

from docx import Document
from pathlib import Path

def verify_company_name(docx_path):
    """Check if old company names still exist in the document."""
    docx_path = Path(docx_path)
    
    if not docx_path.exists():
        print(f"File not found: {docx_path}")
        return
    
    doc = Document(docx_path)
    
    # Collect all text
    all_text = []
    for paragraph in doc.paragraphs:
        all_text.append(paragraph.text)
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    all_text.append(paragraph.text)
    
    full_text = '\n'.join(all_text)
    
    # Check for old names
    old_names = ["بان", "Bany", "bany"]
    new_names = ["لبنداري", "Elbendary", "elbendary"]
    
    print("Verification Results:")
    print("=" * 50)
    
    found_old = False
    found_new = False
    
    for old_name in old_names:
        count = full_text.count(old_name)
        if count > 0:
            print(f"⚠️  Found '{old_name}' {count} time(s) - NOT REPLACED")
            found_old = True
    
    for new_name in new_names:
        count = full_text.count(new_name)
        if count > 0:
            print(f"✓  Found '{new_name}' {count} time(s) - REPLACED")
            found_new = True
    
    if not found_old and found_new:
        print("\n✅ SUCCESS: All old names have been replaced with new names!")
    elif found_old:
        print("\n❌ WARNING: Some old names still exist in the document.")
    else:
        print("\n⚠️  No company names found in the document.")

if __name__ == "__main__":
    doc_path = Path(__file__).parent.parent / "ai" / "rag" / "data" / "bany_developer_profile.docx"
    verify_company_name(doc_path)

